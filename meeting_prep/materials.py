"""Attached files: validate, copy to data/attachments/<meeting_id>/, extract text."""

from __future__ import annotations

import os
import shutil
import zipfile
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
SUPPORTED = (".pdf", ".docx", ".txt", ".md")
MAX_FILES = 10
MAX_TOTAL_BYTES = 20 * 1024 * 1024
MAX_CHARS = 12_000  # per file sent to the model; the brief says when a file was cut


class AttachmentError(ValueError):
    pass


def attachments_root() -> Path:
    return Path(os.environ.get("MEETING_PREP_ATTACHMENTS_DIR") or ROOT / "data" / "attachments")


def _mb(n: int) -> str:
    return f"{n / 1024 / 1024:.1f} MB"


def _check_limits(sizes: dict[str, int], where: str) -> None:
    if len(sizes) > MAX_FILES:
        raise AttachmentError(f"{len(sizes)} files {where}; the limit is {MAX_FILES}. Attach fewer files.")
    total = sum(sizes.values())
    if total > MAX_TOTAL_BYTES:
        raise AttachmentError(f"files {where} total {_mb(total)}; the limit is {_mb(MAX_TOTAL_BYTES)}. "
                              "Attach fewer or smaller files.")


def validate(paths: list[Path]) -> None:
    """Fail before any work if a file is missing, of an unsupported type, or over the limits."""
    for p in paths:
        if not p.is_file():
            raise AttachmentError(f"file not found: {p}")
        if p.suffix.lower() not in SUPPORTED:
            raise AttachmentError(f"unsupported file type: {p.name}. Supported types: {', '.join(SUPPORTED)}.")
    _check_limits({p.name: p.stat().st_size for p in paths}, "attached")


def stage(paths: list[Path], meeting_id: str, root: Path) -> list[Path]:
    """Copy new files in, then return every supported file for the meeting (re-runs reuse earlier ones).
    Limits apply to the combined set and are checked before anything is copied."""
    folder = root / meeting_id
    existing = {p.name: p.stat().st_size for p in _supported(folder)}
    _check_limits(existing | {p.name: p.stat().st_size for p in paths}, f"for {meeting_id} (including {folder})")
    if paths:
        folder.mkdir(parents=True, exist_ok=True)
        for p in paths:
            if p.resolve() != (folder / p.name).resolve():
                shutil.copyfile(p, folder / p.name)
    return _supported(folder)


def _supported(folder: Path) -> list[Path]:
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED)


def _pdf_text(path: Path) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableFile("password-protected")
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except UnreadableFile:
        raise
    except (PdfReadError, ValueError, KeyError, OSError) as exc:
        raise UnreadableFile("corrupt or unreadable PDF") from exc


def _docx_text(path: Path) -> str:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        doc = docx.Document(str(path))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        # Password-protected .docx files are not zip packages, so they land here too.
        raise UnreadableFile("corrupt or password-protected") from exc
    lines = [p.text for p in doc.paragraphs]
    lines += [" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
    return "\n".join(lines)


class UnreadableFile(Exception):
    pass


def read(path: Path) -> dict[str, Any]:
    entry: dict[str, Any] = {"file": path.name, "status": "ok", "reason": None, "text": "", "truncated": False}
    try:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = _pdf_text(path)
        elif suffix == ".docx":
            text = _docx_text(path)
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
        text = text.strip()
        if not text:
            raise UnreadableFile("no readable text")
    except UnreadableFile as exc:
        return entry | {"status": "unreadable", "reason": str(exc)}
    return entry | {"text": text[:MAX_CHARS], "chars": min(len(text), MAX_CHARS), "truncated": len(text) > MAX_CHARS}


def load(paths: list[Path]) -> list[dict[str, Any]]:
    return [read(p) | {"id": f"F{n}"} for n, p in enumerate(paths, 1)]
